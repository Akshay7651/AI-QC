"""Orchestrator tests: agents pool, progress.json, autosave, Ctrl-C, resume, crash handling. Engines are faked."""
import json
import os
import threading
import time
from pathlib import Path

import pandas as pd
import pytest

import local_engine
import progress
import run_qc
from conftest import PNG, make_rows

D = [str(2000000000 + i) for i in range(6)]   # make_rows dockets end in 0..5


def media(tmp_path, dockets=None, with_photos=True):
    d = tmp_path / "media"
    d.mkdir(exist_ok=True)
    for dk in dockets or D:
        (d / f"{dk}.jpg").write_bytes(PNG)            # the form (named by docket)
        if with_photos:
            (d / f"{dk}_1.jpg").write_bytes(PNG)
            (d / f"{dk}_2.jpg").write_bytes(PNG)
    return str(d)


def xlsx(tmp_path, n=6, name="in.xlsx", **fix):
    rows = make_rows(n)
    p = tmp_path / name
    pd.DataFrame(rows).to_excel(p, header=False, index=False)
    return str(p)


def argv(inp, *extra, out="out/o.xlsx"):
    return ["--input", inp, "--output", out, "--checkpoint", "ck/ck.json", "--progress", "out/progress.json", "--no-serve",
            "--engine", "local", "--engine-module", "fakes_engine", "--rate", "0", *extra]


@pytest.fixture(autouse=True)
def _fake_env(tmp_path, monkeypatch):
    monkeypatch.setenv("FAKE_MARK_DIR", str(tmp_path))
    monkeypatch.delenv("FAKE_SLEEP", raising=False)


def read(out="out/o.xlsx"):
    return pd.read_excel(out, dtype=str)


def test_inline_full_run_columns_and_verdicts(tmp_path):
    inp = xlsx(tmp_path)
    assert run_qc.main(argv(inp, "--local-media", media(tmp_path), "--inline", "--agents", "3", "--mode", "full", "--no-risk")) == 0
    df = read()
    assert len(df) == 6
    for c in ["Form No", "PO ID matches docket", "Affected area% (Form)", "Crop Loss% (Form)", "Form vs App (Match/Mismatch/NA)",
              "Match/Mismatch (Form&app)", "Surveyor Signature (Yes/No)", "Farmer Signature (Yes/No)",
              "Primary Worker Signature (Yes/No)", "Government Signature (Yes/No)", "Field photo type",
              "Photo is form image (Yes/No)", "Photo GPS distance (m)", "QC Verdict", "AI_Confidence", "AI_Flags", "Any Other Remarks",
              "docket_id", "farmer_name", "latitude", "Suggested_Remark", "Data_QC_Flags"]:
        assert c in df.columns, c
    by = df.set_index("docket_id")
    assert by.loc[D[0], "QC Verdict"] == "OK" and by.loc[D[0], "Form No"] == "HR0126000000"
    assert by.loc[D[0], "Form vs App (Match/Mismatch/NA)"] == "Match"
    assert by.loc[D[5], "QC Verdict"] == "Reject-evidence" and by.loc[D[5], "Photo is form image (Yes/No)"] == "Yes"
    assert by.loc[D[0], "Primary Worker Signature (Yes/No)"] == "Yes"
    assert "FORM:" in by.loc[D[0], "Any Other Remarks"] and "PHOTOS:" in by.loc[D[0], "Any Other Remarks"]
    assert by.loc[D[0], "Photo GPS distance (m)"] == "12"
    # input columns come first, QC block last, nothing lost
    assert list(df.columns)[:3] == ["docket_id", "application_no", "farmer_name"]
    assert list(df.columns).index("QC Verdict") > list(df.columns).index("Suggested_Remark")


def test_pool_spawn_processes_with_roles(tmp_path):
    inp = xlsx(tmp_path, n=6)
    t0 = time.time()
    assert run_qc.main(argv(inp, "--local-media", media(tmp_path), "--agents", "2", "--no-risk")) == 0
    df = read()
    assert set(df["QC Verdict"]) <= {"OK", "Reject-evidence", "Review"} and df["QC Verdict"].notna().all()
    p = json.load(open("out/progress.json"))
    roles = {a["role"] for a in p["agents"]}
    assert {"Form reader #1", "Photo analyst #1"} <= roles and any(r.startswith("Downloader") for r in roles)
    assert p["status"] == "Done" and p["done"] == p["total"] == 6 and p["percent"] == 100
    assert sum(a["rows_done"] for a in p["agents"] if a["role"].startswith("Form")) + \
        sum(a["rows_done"] for a in p["agents"] if a["role"].startswith("Photo")) >= 12


def test_progress_json_schema(tmp_path):
    run_qc.main(argv(xlsx(tmp_path), "--local-media", media(tmp_path), "--inline", "--no-risk"))
    p = json.load(open("out/progress.json"))
    for k in ["schema", "title", "status", "started_at", "updated_at", "elapsed_sec", "total", "done", "percent", "rows_per_min",
              "eta_sec", "agents", "counters", "verdicts", "events", "errors", "cost", "output"]:
        assert k in p, k
    assert p["cost"] == 0 and len(p["events"]) <= 50 and p["events"][0]["docket"] in D and "remark" in p["events"][0]
    for k in ["form_not_found", "missing_form_link", "photo_is_form", "gps_mismatch", "signature_missing", "mismatch", "overwrite"]:
        assert k in p["counters"]
    assert p["counters"]["photo_is_form"] == 1 and p["counters"]["signature_missing"] == 0
    a = p["agents"][0]
    assert {"id", "role", "state", "docket", "rows_done", "avg_sec"} <= set(a)
    assert sum(p["verdicts"].values()) == 6 and p["output"]["path"].endswith("o.xlsx") and p["output"]["last_saved"]


def test_progress_unit_rate_eta_and_atomic_file(tmp_path):
    pr = progress.Progress(str(tmp_path / "p.json"), total=100)
    pr.add_agent("a1", "Form reader #1")
    pr.agent("a1", state="busy", docket="X")
    pr.agent("a1", state="idle", finished_secs=2.0)
    for i in range(10):
        pr.row_done(f"D{i}", "OK", "remark", ["mismatch"])
    s = pr.snapshot()
    assert s["done"] == 10 and s["counters"]["mismatch"] == 10 and s["agents"][0]["avg_sec"] == 2.0
    assert s["eta_sec"] is not None and s["percent"] == 10.0
    for i in range(80):
        pr.row_done(f"E{i}", "OK", "r")
    assert len(pr.snapshot()["events"]) == 50
    pr.write_now()
    assert json.loads((tmp_path / "p.json").read_text())["done"] == 90
    assert not list(tmp_path.glob("*.tmp"))


def test_http_server_serves_dashboard_and_json(tmp_path):
    import urllib.request
    pr = progress.Progress(str(tmp_path / "p.json"), total=5)
    pr.start(serve_port=18765, host="127.0.0.1")
    try:
        assert pr.url
        html = urllib.request.urlopen(pr.url, timeout=5).read().decode()
        assert "progress.json" in html and "<title>" in html
        d = json.loads(urllib.request.urlopen(pr.url + "progress.json", timeout=5).read())
        assert d["total"] == 5
        with pytest.raises(Exception):
            urllib.request.urlopen(pr.url + "secret.txt", timeout=5)
    finally:
        pr.stop("Done")


def test_one_bad_row_does_not_kill_run_and_dead_worker_is_retried(tmp_path, monkeypatch):
    monkeypatch.setenv("FAKE_DEATH", "1")
    # dockets ending 9 (engine raises), 8 (worker dies once -> retry succeeds), 4 (dies every time -> error row)
    rows = make_rows(6)
    for r, last in zip(rows, ["9", "8", "4", "0", "1", "2"]):
        r[0] = f"200000000{last}"
    p = tmp_path / "bad.xlsx"
    pd.DataFrame(rows).to_excel(p, header=False, index=False)
    ds = [r[0] for r in rows]
    assert run_qc.main(argv(str(p), "--local-media", media(tmp_path, ds), "--agents", "2", "--no-risk", "--task-timeout", "60")) == 0
    by = read().set_index("docket_id")
    assert len(by) == 6
    assert by.loc["2000000009", "QC Verdict"] == "Manual-check" and "fake engine exploded" in by.loc["2000000009", "Any Other Remarks"] or \
        by.loc["2000000009", "QC Verdict"] in ("Manual-check", "Reject-evidence")
    assert by.loc["2000000008", "Form No"] == "HR0126000008"          # died once, retried OK
    assert by.loc["2000000004", "Form No"] != by.loc["2000000004", "Form No"] or pd.isna(by.loc["2000000004", "Form No"])
    assert by.loc["2000000004", "QC Verdict"] in ("Manual-check", "Reject-evidence", "Review")
    assert by.loc["2000000000", "QC Verdict"] == "OK"
    pj = json.load(open("out/progress.json"))
    assert pj["error_count"] >= 3 and pj["done"] == 6          # worker deaths are visible on the dashboard


def test_no_media_rows_get_missing_link_remarks(tmp_path):
    rows = make_rows(3)
    for r in rows:
        r[24], r[25] = None, None
    p = tmp_path / "nolinks.xlsx"
    pd.DataFrame(rows).to_excel(p, header=False, index=False)
    assert run_qc.main(argv(str(p), "--inline", "--no-risk")) == 0
    df = read()
    assert set(df["QC Verdict"]) == {"Reject-evidence"}
    assert df["Any Other Remarks"].str.contains("No signed-form link").all() and df["Any Other Remarks"].str.contains("No photo link").all()
    assert json.load(open("out/progress.json"))["counters"]["missing_form_link"] == 3


def test_unreachable_urls_are_manual_check_not_fatal(tmp_path):
    assert run_qc.main(argv(xlsx(tmp_path), "--inline", "--no-risk")) == 0     # conftest: every download answers HTTP 404
    df = read()
    assert set(df["QC Verdict"]) == {"Manual-check"} and df["Any Other Remarks"].str.contains("HTTP 404").all()
    assert json.load(open("out/progress.json"))["counters"]["form_not_found"] == 6


def test_autosave_writes_partial_file_while_running(tmp_path, monkeypatch):
    monkeypatch.setenv("FAKE_SLEEP", "0.5")
    inp = xlsx(tmp_path)
    m = media(tmp_path)
    res = {}
    th = threading.Thread(target=lambda: res.update(rc=run_qc.main(
        argv(inp, "--local-media", m, "--inline", "--agents", "1", "--autosave-sec", "0.7", "--no-risk"))))
    th.start()
    seen = None
    t_end = time.time() + 30
    while time.time() < t_end and th.is_alive():
        if os.path.exists("out/o.xlsx"):
            try:
                df = pd.read_excel("out/o.xlsx", dtype=str)
            except Exception:
                continue
            if 0 < df["QC Verdict"].notna().sum() < 6:
                seen = df
                break
        time.sleep(0.1)
    th.join(60)
    assert seen is not None, "no partial autosave observed while the run was in progress"
    assert len(seen) == 6 and seen["docket_id"].notna().all()               # unprocessed rows are still present
    assert res["rc"] == 0 and read()["QC Verdict"].notna().all()
    assert json.load(open("ck/ck.json"))["results"]


def test_ctrl_c_saves_partial_and_resume_finishes(tmp_path, monkeypatch):
    inp = xlsx(tmp_path)
    m = media(tmp_path)
    calls = {"n": 0}
    real_finalize = local_engine.LocalRunner._finalize

    def finalize_then_interrupt(self, i):
        real_finalize(self, i)
        calls["n"] += 1
        if calls["n"] == 2:
            raise KeyboardInterrupt

    monkeypatch.setattr(local_engine.LocalRunner, "_finalize", finalize_then_interrupt)
    rc = run_qc.main(argv(inp, "--local-media", m, "--inline", "--agents", "1", "--no-risk"))
    assert rc == 130
    monkeypatch.setattr(local_engine.LocalRunner, "_finalize", real_finalize)
    part = read()
    assert len(part) == 6 and 1 <= part["QC Verdict"].notna().sum() < 6
    assert json.load(open("ck/ck.json"))["results"]
    assert json.load(open("out/progress.json"))["status"] in ("Stopped", "Done")
    # resume only does the remaining rows
    import fakes_engine
    fakes_engine.CALLS["form"] = 0
    assert run_qc.main(argv(inp, "--local-media", m, "--inline", "--agents", "1", "--no-risk", "--resume")) == 0
    assert read()["QC Verdict"].notna().all()
    assert fakes_engine.CALLS["form"] < 6


def test_resume_reuses_checkpoint_and_skips_done_rows(tmp_path):
    import fakes_engine
    inp, m = xlsx(tmp_path), media(tmp_path)
    run_qc.main(argv(inp, "--local-media", m, "--inline", "--no-risk"))
    fakes_engine.CALLS.update(form=0, photo=0)
    run_qc.main(argv(inp, "--local-media", m, "--inline", "--no-risk", "--resume"))
    assert fakes_engine.CALLS == {"form": 0, "photo": 0}
    assert read()["QC Verdict"].notna().all()


def test_excel_locked_falls_back_to_timestamped_sibling(tmp_path, monkeypatch, capsys):
    real = os.replace
    target = os.path.abspath("out/o.xlsx")

    def locked(src, dst):
        if os.path.abspath(dst) == target:
            raise PermissionError("file is open in Excel")
        return real(src, dst)
    monkeypatch.setattr(os, "replace", locked)
    rc = run_qc.main(argv(xlsx(tmp_path), "--local-media", media(tmp_path), "--inline", "--no-risk"))
    assert rc == 0
    files = [p.name for p in Path("out").glob("o_*.xlsx")]
    assert len(files) == 1 and not Path("out/o.xlsx").exists()
    assert "open/locked" in capsys.readouterr().out
    assert not list(Path("out").glob("*.tmp*"))


def test_never_overwrites_input_unless_inplace(tmp_path):
    inp = xlsx(tmp_path)
    with pytest.raises(SystemExit):
        run_qc.main(["--input", inp, "--output", inp, "--mode", "gps", "--no-serve"])
    before = Path(inp).read_bytes()
    assert run_qc.main(["--input", inp, "--mode", "gps", "--no-serve", "--checkpoint", "ck/c.json"]) == 0
    assert Path(inp).read_bytes() == before and Path(str(Path(inp).with_suffix("")) + "_QC.xlsx").exists()
    assert run_qc.main(["--input", inp, "--inplace", "--mode", "gps", "--no-serve", "--checkpoint", "ck/c.json"]) == 0
    assert "Suggested_Remark" in pd.read_excel(inp).columns


def test_agents_alias_default_and_clamp():
    a = run_qc.parse_args(["--input", "x.xlsx", "--agents", "12"])
    assert a.workers == 12 and a.serve_port == 8765 and a.autosave_sec == 60 and not a.no_serve
    assert run_qc.parse_args(["--input", "x.xlsx", "--workers", "3"]).workers == 3
    assert run_qc.parse_args(["--input", "x.xlsx"]).workers is None


def test_excel_is_readable_format_and_keeps_existing_human_remark(tmp_path):
    from openpyxl import load_workbook
    rows = make_rows(2)
    df = pd.DataFrame(rows)
    p = tmp_path / "in.xlsx"
    df.to_excel(p, header=False, index=False)
    run_qc.main(argv(str(p), "--local-media", media(tmp_path), "--inline", "--no-risk"))
    wb = load_workbook("out/o.xlsx")
    ws = wb.active
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref
    hdr = [c.value for c in ws[1]]
    assert "Any Other Remarks" in hdr and ws.conditional_formatting


def test_classify_local_splits_forms_and_photos(tmp_path):
    f = {"pdfs": [Path("a/D1.pdf")], "images": [Path("a/D1.jpg"), Path("a/D1_1.jpg"), Path("a/D1_2.jpg"), Path("forms/x.jpg")]}
    forms, photos = local_engine.classify_local(f, "D1")
    assert [p.name for p in forms][0] == "D1.pdf" and {p.name for p in forms} == {"D1.pdf", "D1.jpg", "x.jpg"}
    assert [p.name for p in photos] == ["D1_1.jpg", "D1_2.jpg"]


def test_assemble_local_keeps_input_values_for_unprocessed_rows():
    import run_qc as r
    from helpers import frame, good
    df = frame([good(0, docket_id="A"), good(1, docket_id="B")])
    df["Any Other Remarks"] = ["human note", "kept"]
    df["Done By"] = ["Human", "Human"]
    res = {"A": {"verdict": "OK", "confidence": "High", "flags": [], "remark": "OK: fine", "match": "Match", "form_area": 5.0,
                 "form_loss": 5.0, "form": {"_state": "ok"}, "photo": {"_state": "skipped"}, "n_photos": 0, "ts": "t", "engine": "local"}}
    out = r.assemble_local(df, res, ["A", "B"])
    assert out["Done By"].tolist() == ["AI-Auto", "Human"]
    assert out["Any Other Remarks"][0] == "OK: fine | INPUT REMARK: human note" and out["Any Other Remarks"][1] == "kept"
    assert out["QC Verdict"].tolist()[0] == "OK" and pd.isna(out["QC Verdict"][1]) and out["AI_Flags"][1] == ""
    # an earlier AI remark is replaced, not stacked
    df["Any Other Remarks"] = ["OK: old | FORM: x", "kept"]
    assert r.assemble_local(df, res, ["A", "B"])["Any Other Remarks"][0] == "OK: fine"
