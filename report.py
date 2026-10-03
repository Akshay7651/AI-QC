"""Summary workbook (summary_report.xlsx): ONE sheet 'Summary' with the run's numbers, built from the QC output columns."""
import pandas as pd

import config as C

VERDICTS = ["OK", "Partially OK", "Review", "Manual QC Required"]


def _col(df, *names):
    for n in names:
        if n in df:
            return df[n]
    return pd.Series([None] * len(df), index=df.index)


def _field_counts(s):
    """value column -> (read, can't read, blank on form, no form / not checked)"""
    v = s.fillna("").astype(str).str.strip()
    cant = v.str.lower().str.startswith(("mentioned but can't read", "can't read", "cant read"))
    blank = v.str.lower().str.startswith("blank")
    none = v.eq("") | v.str.lower().isin(["na", "no form", "nan"])
    read = ~(cant | blank | none)
    return int(read.sum()), int(cant.sum()), int(blank.sum()), int(none.sum())


def _tags(remark):
    """'Manual QC Required - A, B, C' -> ['A', 'B', 'C']"""
    r = str(remark or "")
    if " - " not in r:
        return []
    return [t.strip() for t in r.split(" - ", 1)[1].split(",") if t.strip()]


def build(df: pd.DataFrame, path: str):
    n = len(df)
    verdict = _col(df, "QC Verdict").fillna("")
    match = _col(df, C.COL_MATCH).fillna("NA")
    remark = _col(df, "AI Remark", "Suggested_Remark").fillna("")
    dist = _col(df, "district").fillna("Unknown")
    surv = _col(df, "surveyor_name").fillna("Unknown")
    tags = remark.map(_tags)

    blocks = []          # (title, DataFrame)
    pct = lambda k: round(100.0 * k / n, 1) if n else 0.0
    ov = [("Rows checked", n, 100.0 if n else 0.0)] + [(v, int((verdict == v).sum()), pct(int((verdict == v).sum()))) for v in VERDICTS]
    ov += [("Form vs app: Match", int((match == "Match").sum()), pct(int((match == "Match").sum()))),
           ("Form vs app: Mismatch", int((match == "Mismatch").sum()), pct(int((match == "Mismatch").sum()))),
           ("Form vs app: not comparable (value not read)", int((~match.isin(["Match", "Mismatch"])).sum()), pct(int((~match.isin(["Match", "Mismatch"])).sum())))]
    blocks.append(("1. Overall result", pd.DataFrame(ov, columns=["Item", "Rows", "% of rows"])))

    fr = []
    for label, cols in (("Form No", ("Form No",)), ("PO ID (Form)", ("PO ID (Form)",)),
                        ("Affected area % (Form)", (C.COL_FORM_AREA, "Affected area% (Form)")),
                        ("Crop loss % (Form)", (C.COL_FORM_LOSS, "Crop Loss% (Form)"))):
        r, c, b, x = _field_counts(_col(df, *cols))
        fr.append((label, r, pct(r), c, b, x))
    po, dk = _col(df, "PO ID (Form)").fillna("").astype(str), _col(df, "docket_id").fillna("").astype(str)
    fr.append(("PO ID equal to docket", int((po == dk).sum()), pct(int((po == dk).sum())), None, None, None))
    for label, cols in (("Farmer signature = Yes", ("Farmer Signature (Yes/No)",)), ("Surveyor signature = Yes", ("Surveyor Signature (Yes/No)",)),
                        ("Government signature = Yes", ("Government Signature (Yes/No)",)), ("Field photo = Yes", (C.COL_FIELD_PHOTO, "Field photo (no crop / cut & spread / crop mismatch / standing crop)")),
                        ("Farmer photo = Yes", ("Farmer Photo (Yes/No)",)), ("Survey remarks on form = Yes", ("Survey remarks on form",))):
        k = int(_col(df, *cols).fillna("").astype(str).str.strip().str.lower().eq("yes").sum())
        fr.append((label, k, pct(k), None, None, None))
    blocks.append(("2. What the AI could read on the forms", pd.DataFrame(fr, columns=["Field", "Read", "% of rows", "Written but can't read", "Blank on form", "No form / not checked"])))

    cnt = {}
    for ts in tags:
        for t in ts:
            cnt[t] = cnt.get(t, 0) + 1
    iss = sorted(cnt.items(), key=lambda kv: -kv[1])
    blocks.append(("3. Issues found (a row can have several)", pd.DataFrame([(t, k, pct(k)) for t, k in iss], columns=["Issue", "Rows", "% of rows"])))

    def per(group, name):
        t = pd.DataFrame({name: group, "Rows": 1, "Match": match.eq("Match"), "Mismatch": match.eq("Mismatch"),
                          "Same spot": tags.map(lambda x: any("Same" in y for y in x)).astype(bool)})
        for v in VERDICTS:
            t[v] = verdict.eq(v)
        g = t.groupby(name).sum(numeric_only=True).reset_index()
        g["OK %"] = (100 * g["OK"] / g["Rows"]).round(1)
        g["Match % (of compared)"] = (100 * g["Match"] / (g["Match"] + g["Mismatch"]).replace(0, float("nan"))).round(1)
        return g[[name, "Rows"] + VERDICTS + ["OK %", "Match", "Mismatch", "Match % (of compared)", "Same spot"]].sort_values("Rows", ascending=False)
    blocks.append(("4. By district", per(dist, "District")))
    blocks.append(("5. By surveyor", per(surv, "Surveyor")))

    with pd.ExcelWriter(path, engine="openpyxl") as w:
        r0 = 0
        for title, t in blocks:
            pd.DataFrame([[title]]).to_excel(w, sheet_name="Summary", startrow=r0, index=False, header=False)
            t.to_excel(w, sheet_name="Summary", startrow=r0 + 1, index=False)
            r0 += len(t) + 4
        ws = w.sheets["Summary"]
        from openpyxl.styles import Font, PatternFill
        head = PatternFill("solid", fgColor="1F3A5F")
        for row in ws.iter_rows():
            for c in row:
                if c.column == 1 and isinstance(c.value, str) and c.value[:2] in {f"{i}." for i in range(1, 10)}:
                    c.font = Font(bold=True, size=13, color="1F3A5F")
        r0 = 0
        for title, t in blocks:
            for j in range(len(t.columns)):
                cell = ws.cell(row=r0 + 2, column=j + 1)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
            r0 += len(t) + 4
        ws.column_dimensions["A"].width = 44
        for col in "BCDEFGHIJKLM":
            ws.column_dimensions[col].width = 16


# ============================================================================ the big QC workbook (fast, atomic)
import os
import time
from datetime import datetime
from pathlib import Path

VERDICT_COLOURS = {"Manual QC Required": "#F8BBD0", "Review": "#FFE0B2", "Partially OK": "#FFF9C4", "OK": "#C8E6C9"}
CONF_COLOURS = {"Low": "#F8BBD0", "Medium": "#FFF9C4", "High": "#C8E6C9"}
BAD = "#F8BBD0"


from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as _ILLEGAL


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:
        return None
    if v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.strftime("%Y-%m-%d %H:%M:%S") if v == v else None
    if isinstance(v, str):
        return _ILLEGAL.sub("", v) if _ILLEGAL.search(v) else v
    if isinstance(v, (bool, int, float)):
        return v
    if hasattr(v, "item"):
        try:
            return _cell_value(v.item())
        except Exception:
            pass
    return str(v)


def write_xlsx_slow(df: pd.DataFrame, path: str, remarks_col="Any Other Remarks"):
    """One sheet, frozen header + docket column, widths, wrapped remarks, colour rules. xlsxwriter (constant memory) if present,
    else openpyxl write-only. Returns the writer name."""
    cols = [str(c) for c in df.columns]
    try:
        import xlsxwriter
    except ImportError:
        return _write_openpyxl(df, path, cols)
    wb = xlsxwriter.Workbook(path, {"constant_memory": True, "nan_inf_to_errors": True, "strings_to_formulas": False,
                                    "strings_to_urls": False})
    ws = wb.add_worksheet("QC")
    head = wb.add_format({"bold": True, "bg_color": "#1F3A5F", "font_color": "#FFFFFF", "text_wrap": True, "valign": "top", "border": 1})
    wrap = wb.add_format({"text_wrap": True, "valign": "top"})
    sample = df.head(300)
    widths = []
    for c in cols:
        s = sample[c].astype(str).str.len() if len(sample) else pd.Series([0])
        w = min(60, max(len(c) * 0.9, float(s.quantile(0.9)) if len(s) else 8) + 2)
        widths.append(max(8, w))
    wrap_cols = {cols.index(remarks_col)} if remarks_col in cols else set()
    for j, c in enumerate(cols):
        if j in wrap_cols:
            ws.set_column(j, j, 110, wrap)
        else:
            ws.set_column(j, j, widths[j])
    ws.set_row(0, 45)
    ws.write_row(0, 0, cols, head)
    arr = df.astype(object).to_numpy()
    for r in range(arr.shape[0]):
        row = arr[r]
        for j in range(len(cols)):
            v = _cell_value(row[j])
            if v is None:
                continue
            if isinstance(v, str):
                ws.write_string(r + 1, j, v)
            elif isinstance(v, bool):
                ws.write_boolean(r + 1, j, v)
            else:
                ws.write_number(r + 1, j, v)
    n = len(df) + 1
    ws.freeze_panes(1, 1 if cols and cols[0].lower().startswith(("docket", "application")) else 0)
    if len(cols):
        ws.autofilter(0, 0, max(n - 1, 1), len(cols) - 1)

    def cf(col, criteria, colour):
        if col in cols:
            j = cols.index(col)
            ws.conditional_format(1, j, max(n, 2), j, {"type": "cell", "criteria": "==", "value": f'"{criteria}"',
                                                        "format": wb.add_format({"bg_color": colour})})
    for k, colour in VERDICT_COLOURS.items():
        cf("QC Verdict", k, colour)
    for k, colour in CONF_COLOURS.items():
        cf("AI_Confidence", k, colour)
    for c in cols:
        if c.startswith("Match/Mismatch") or c.startswith("Form vs App"):
            cf(c, "Mismatch", BAD)
            cf(c, "Match", "#C8E6C9")
        elif "Signature (Yes/No)" in c:
            cf(c, "No", BAD)
        elif c.startswith("Photo is form image"):
            cf(c, "Yes", BAD)
        elif c.startswith("PO ID matches"):
            cf(c, "No", BAD)
        elif c.startswith("Form Status"):
            cf(c, "overwrite", BAD)
            cf(c, "incomplete", "#FFF9C4")
    wb.close()
    return "xlsxwriter"


def _write_openpyxl(df, path, cols):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("QC")
    ws.freeze_panes = "A2"
    ws.append(cols)
    for row in df.astype(object).to_numpy():
        out = []
        for v in row:
            v = _cell_value(v)
            out.append(ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)
        ws.append(out)
    wb.save(path)
    return "openpyxl"


class ExcelTarget:
    """Atomic writes (tmp + replace). If the file is locked (Excel open on Windows) fall back to a timestamped sibling."""

    def __init__(self, path):
        self.path = str(path)
        self.fallback = None
        self.last_note = ""

    def write(self, df):
        Path(self.path).parent.mkdir(parents=True, exist_ok=True)
        tmp = f"{self.path}.tmp{os.getpid()}.xlsx"
        try:
            write_xlsx(df, tmp)
            try:
                os.replace(tmp, self.path)
                self.last_note = ""
                return self.path
            except PermissionError:
                fb = self.fallback or str(Path(self.path).with_name(
                    f"{Path(self.path).stem}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"))
                os.replace(tmp, fb)
                self.fallback = fb
                self.last_note = f"{self.path} is open/locked - saved to {fb} instead"
                return fb
        finally:
            try:
                if os.path.exists(tmp):
                    os.remove(tmp)
            except OSError:
                pass


# ---- fast direct-XML writer (about 10x faster than xlsxwriter for 40k x 80 cells) -------------------------------
import re as _re
import zipfile as _zip
from xml.sax.saxutils import escape as _esc

_CTRL = _re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _col_letters(n):
    s = ""
    n += 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


AI_COLS = set()      # names of the columns written by the AI (set by run_qc); they get the green colour family
_DXF = {"red": "F8BBD0", "amber": "FFE0B2", "yellow": "FFF9C4", "green": "C8E6C9"}


def write_xlsx(df: pd.DataFrame, path: str, remarks_col="Any Other Remarks"):
    """One sheet 'QC': frozen header+first id column, widths, wrapped remarks, colour rules, autofilter. Streams XML straight
    into the zip (no per-cell library overhead). Falls back to xlsxwriter/openpyxl if anything goes wrong."""
    try:
        return _write_fast(df, path, remarks_col)
    except Exception:
        return write_xlsx_slow(df, path, remarks_col)


def _write_fast(df, path, remarks_col):
    cols = [str(c) for c in df.columns]
    nc, nr = len(cols), len(df)
    L = [_col_letters(j) for j in range(nc)]
    sample = df.head(300)
    widths = []
    for c in cols:
        s = sample[c].astype(str).str.len() if nr else pd.Series([0])
        widths.append(max(8, min(60, max(len(c) * 0.9, float(s.quantile(0.9)) if len(s) else 8) + 2)))
    rem_j = cols.index(remarks_col) if remarks_col in cols else -1
    # two colour families: the original (raw) columns are blue, the columns written by the AI are green
    is_ai = [c in AI_COLS for c in cols]
    body = [(5 if (j == rem_j or cols[j] == "AI Technical Detail") else 4) if is_ai[j] else 3 for j in range(nc)]
    SA = [f' s="{x}"' for x in body]
    HS = [6 if is_ai[j] else 1 for j in range(nc)]
    colxml = "".join(f'<col min="{j + 1}" max="{j + 1}" width="{110 if j == rem_j else (80 if cols[j] == "AI Technical Detail" else widths[j]):.1f}" customWidth="1" style="{body[j]}"/>'
                     for j in range(nc))
    # conditional formats: (column index, value, dxf id)
    dxf_ids = {k: i for i, k in enumerate(_DXF)}
    rules = []

    def cf(c, val, colour):
        if c in cols:
            rules.append((cols.index(c), val, dxf_ids[colour]))
    for k, colour in {"Manual QC Required": "red", "Review": "amber", "Partially OK": "yellow", "OK": "green"}.items():
        cf("QC Verdict", k, colour)
    for k, colour in {"Low": "red", "Medium": "yellow", "High": "green"}.items():
        cf("AI_Confidence", k, colour)
    for c in cols:
        if c.startswith("Match/Mismatch") or c.startswith("Form vs App"):
            cf(c, "Mismatch", "red")
            cf(c, "Match", "green")
        elif "Signature (Yes/No)" in c or c.startswith("PO ID matches"):
            cf(c, "No", "red")
        elif c.startswith("Photo is form image"):
            cf(c, "Yes", "red")
        elif c.startswith("Form Status"):
            cf(c, "overwrite", "red")
            cf(c, "incomplete", "yellow")
    last = nr + 1
    cfxml, prio = [], 1
    for j, val, d in rules:
        cfxml.append(f'<conditionalFormatting sqref="{L[j]}2:{L[j]}{max(last, 2)}"><cfRule type="cellIs" dxfId="{d}" priority="{prio}" '
                     f'operator="equal"><formula>"{_esc(val)}"</formula></cfRule></conditionalFormatting>')
        prio += 1
    xsplit = 1 if cols and cols[0].lower().startswith(("docket", "application")) else 0
    pane = (f'<pane xSplit="{xsplit}" ySplit="1" topLeftCell="{"B" if xsplit else "A"}2" activePane="{"bottomRight" if xsplit else "bottomLeft"}" state="frozen"/>')
    head = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetViews><sheetView workbookViewId="0">{pane}</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/>'
            f'<cols>{colxml}</cols><sheetData>')
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font><font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font></fonts>'
              '<fills count="6"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill>'
              '<fill><patternFill patternType="solid"><fgColor rgb="FF1F3A5F"/><bgColor indexed="64"/></patternFill></fill>'
              '<fill><patternFill patternType="solid"><fgColor rgb="FFE3EEFB"/><bgColor indexed="64"/></patternFill></fill>'
              '<fill><patternFill patternType="solid"><fgColor rgb="FFE6F4EA"/><bgColor indexed="64"/></patternFill></fill>'
              '<fill><patternFill patternType="solid"><fgColor rgb="FF1E7B4F"/><bgColor indexed="64"/></patternFill></fill></fills>'
              '<borders count="2"><border><left/><right/><top/><bottom/><diagonal/></border><border><left style="thin"><color rgb="FFB7C0CC"/></left><right style="thin"><color rgb="FFB7C0CC"/></right><top style="thin"><color rgb="FFB7C0CC"/></top><bottom style="thin"><color rgb="FFB7C0CC"/></bottom><diagonal/></border></borders>'
              '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
              '<cellXfs count="7"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
              '<xf numFmtId="0" fontId="1" fillId="2" borderId="1" xfId="0" applyBorder="1" applyFont="1" applyFill="1" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>'
              '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>'
              '<xf numFmtId="0" fontId="0" fillId="3" borderId="1" xfId="0" applyBorder="1" applyFill="1" applyAlignment="1"><alignment vertical="top"/></xf>'
              '<xf numFmtId="0" fontId="0" fillId="4" borderId="1" xfId="0" applyBorder="1" applyFill="1" applyAlignment="1"><alignment vertical="top"/></xf>'
              '<xf numFmtId="0" fontId="0" fillId="4" borderId="1" xfId="0" applyBorder="1" applyFill="1" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>'
              '<xf numFmtId="0" fontId="1" fillId="5" borderId="1" xfId="0" applyBorder="1" applyFont="1" applyFill="1" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf></cellXfs>'
              '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
              f'<dxfs count="{len(_DXF)}">' + "".join(f'<dxf><fill><patternFill patternType="solid"><fgColor rgb="FF{v}"/><bgColor rgb="FF{v}"/></patternFill></fill></dxf>'
                                                    for v in _DXF.values()) + '</dxfs></styleSheet>')
    wbxml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="QC" sheetId="1" r:id="rId1"/></sheets></workbook>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wbrels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
              '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
              '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
              '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
    esc, ctrl = _esc, _CTRL

    def rows():
        yield head
        yield '<row r="1" ht="45" customHeight="1">' + "".join(
            f'<c r="{L[j]}1" s="{HS[j]}" t="inlineStr"><is><t>{esc(c)}</t></is></c>' for j, c in enumerate(cols)) + "</row>"
        arr = df.astype(object).to_numpy()
        buf = []
        for r in range(nr):
            rn = r + 2
            cells = []
            row = arr[r]
            for j in range(nc):
                v = row[j]
                if v is None:
                    continue
                t = type(v)
                if t is str:
                    if not v:
                        continue
                    if ctrl.search(v):
                        v = ctrl.sub("", v)
                    if len(v) > 32000:
                        v = v[:32000]
                    cells.append(f'<c r="{L[j]}{rn}"{SA[j]} t="inlineStr"><is><t xml:space="preserve">{esc(v)}</t></is></c>')
                elif t is bool:
                    cells.append(f'<c r="{L[j]}{rn}"{SA[j]} t="b"><v>{int(v)}</v></c>')
                elif t is int or t is float:
                    if v != v or v in (float("inf"), float("-inf")):
                        continue
                    cells.append(f'<c r="{L[j]}{rn}"{SA[j]}><v>{v!r}</v></c>')
                else:
                    v = _cell_value(v)
                    if v is None:
                        continue
                    if isinstance(v, bool):
                        cells.append(f'<c r="{L[j]}{rn}"{SA[j]} t="b"><v>{int(v)}</v></c>')
                    elif isinstance(v, (int, float)):
                        cells.append(f'<c r="{L[j]}{rn}"{SA[j]}><v>{v!r}</v></c>')
                    else:
                        cells.append(f'<c r="{L[j]}{rn}"{SA[j]} t="inlineStr"><is><t xml:space="preserve">{esc(str(v))}</t></is></c>')
            buf.append(f'<row r="{rn}">' + "".join(cells) + "</row>")
            if len(buf) >= 2000:
                yield "".join(buf)
                buf = []
        if buf:
            yield "".join(buf)
        yield ("</sheetData>" + (f'<autoFilter ref="A1:{L[-1]}{max(last, 2)}"/>' if nc else "") + "".join(cfxml) + "</worksheet>")

    with _zip.ZipFile(path, "w", _zip.ZIP_DEFLATED, compresslevel=3) as z:
        z.writestr("[Content_Types].xml", ctypes)
        z.writestr("_rels/.rels", rels)
        z.writestr("xl/workbook.xml", wbxml)
        z.writestr("xl/_rels/workbook.xml.rels", wbrels)
        z.writestr("xl/styles.xml", styles)
        with z.open("xl/worksheets/sheet1.xml", "w", force_zip64=True) as f:
            for chunk in rows():
                f.write(chunk.encode("utf-8"))
    return "fastxml"
